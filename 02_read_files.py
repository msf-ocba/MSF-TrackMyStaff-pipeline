"""
02_read_files.py
------------------
Finds the SINGLE latest extraction file (by YYMMDD embedded in its
filename) across ALL folders under download_root, and reproduces the
manual pre-processing steps from spec section 4 ("File pre-processing"),
steps 3-6, on that one file only:

  3. Save/rename as A1_EXTRACTION_TMS_OCBA_YYMMDD.csv
  4. Create a file copy, rename it, change extension to .txt:
     A2_EXTRACTION_TMS_OCBA_YYMMDD.txt
  5. "Open this txt file with Excel & Transform it via text import
     wizard: check 'Delimited' > Next > check only 'semicolon' > Next
     > Select column for Serial Numbers (PCL_NO_SERIE_LOT) and choose
     'Text' > Finish"
     -> Replicated by parsing the semicolon-delimited file and forcing
        PCL_NO_SERIE_LOT to text.
  6. Create a table with headers, save as xlsx:
     A3_EXTRACTION_TMS_OCBA_YYMMDD.xlsx

IMPORTANT: MSF Logistique extractions are weekly snapshots
identified by the YYMMDD in their filename, and each week's batch is
processed entirely on its own. So this script:
  - scans every file under download_root (all dated download subfolders)
  - picks the ONE file with the latest YYMMDD in its name
  - processes only that file

work_dir and output_dir are dated per BATCH RUN, using the ISO week
BEFORE the one this script is executed in (e.g. work/Y26W25/ if run
during week 26 - see common_config.compute_week_tag()), not the YYMMDD
embedded in the source filename - this matches the same week-tag
convention used by 04_validate.py and 05_export.py for their own
output folders, so every stage of the pipeline organizes its files the
same way. A1,
A2, A3 (and later A4, A5, A6) filenames still embed the batch's own
YYMMDD date tag, taken from the source filename - only the FOLDER name
changed. Every batch processed in the same calendar week lands in the
same work folder; the A-prefixed filenames (which include the date
tag) are what keep different batches from colliding within that folder.

Source files are semicolon-delimited, double-quoted, ISO-8859-1
("latin1") encoded, with comma decimal separators (e.g. "2081,40") and
DD/MM/YYYY dates - this matches the real MSF Logistique export format.

Usage:
    python 02_read_files.py [path_to_specific_file.csv]

If a specific file path is given, that file is used directly (its
YYMMDD is still parsed from its filename to name the A1/A2/A3 outputs).
If omitted, the latest file by YYMMDD across download_root is
auto-selected.
"""
import os
import re
import sys
import glob
import shutil
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from common_config import load_config, get_paths, get_dated_subdir, compute_week_tag

# Real MSF Logistique extraction header (see spec section 2 + actual
# export files). Keep this in sync if MSF Logistique changes their
# column layout.
EXPECTED_COLUMNS = [
    "FCL_ART_CODE",
    "ART_DES1",
    "PCL_NO_SERIE_LOT",
    "PCT_REF_CMDE1",
    "FCT_NO_FACTURE",
    "FCT_CLI_CODE_FAC",
    "PCT_CLI_CODE_LIV",
    "PRIX_VENTE_UNIT",
    "QTE",
    "SDT_DT_ENLEV",
    "CODE_KIT",
    "KIT_NO_SERIE_LOT",
    "KIT_LOT_NO",
    "DATE_FABRICATION",
    "GARANTIE",
    "MARQUE",
    "MODEL",
]

# Columns that must always be read as text, never inferred as numbers,
# because they can contain leading zeros, letters, or are pure identifiers.
# PCL_NO_SERIE_LOT is the one the manual process explicitly calls out
# in the Text Import Wizard step.
TEXT_COLUMNS = [
    "FCL_ART_CODE",
    "PCL_NO_SERIE_LOT",
    "PCT_REF_CMDE1",
    "FCT_NO_FACTURE",
    "FCT_CLI_CODE_FAC",
    "PCT_CLI_CODE_LIV",
    "CODE_KIT",
    "KIT_NO_SERIE_LOT",
    "KIT_LOT_NO",
]

# Matches a 6-digit YYMMDD immediately before the file extension, e.g.
# "EXTRACTION_TMS_OCBA_260623.csv" or "A1_EXTRACTION_TMS_OCBA_260623.csv".
DATE_TAG_RE = re.compile(r"(\d{6})(?=\.\w+$)")


def extract_date_tag(filename):
    """Pull the YYMMDD date tag out of a filename. Returns None if absent."""
    match = DATE_TAG_RE.search(os.path.basename(filename))
    return match.group(1) if match else None


def find_latest_file_by_date(download_root):
    """
    Scan every file under download_root (any depth, any dated
    subfolder) and return the path of the file whose filename contains
    the latest YYMMDD tag. Only files with a parseable date tag are
    considered.
    """
    all_files = glob.glob(os.path.join(download_root, "**", "*"), recursive=True)
    candidates = []
    for path in all_files:
        if not os.path.isfile(path):
            continue
        tag = extract_date_tag(path)
        if tag:
            candidates.append((tag, path))

    if not candidates:
        raise FileNotFoundError(
            f"No files with a YYMMDD date tag found anywhere under '{download_root}'. "
            f"Run 01_download_sftp.py first."
        )

    candidates.sort(key=lambda t: t[0])  # YYMMDD sorts correctly as a string
    latest_tag, latest_path = candidates[-1]
    return latest_path, latest_tag


def read_one_csv(path):
    dtype_map = {col: str for col in TEXT_COLUMNS}
    df = pd.read_csv(
        path,
        sep=";",
        dtype=dtype_map,
        encoding="latin1",   # MSF Logistique exports are ISO-8859-1 / Western European
        keep_default_na=True,
        quotechar='"',
    )
    df["SOURCE_FILE"] = os.path.basename(path)
    return df


def step3_save_a1(source_path, work_dir, date_tag):
    """Spec step 3: save the single source file as A1_..._YYMMDD.csv"""
    df = read_one_csv(source_path)

    missing = [c for c in EXPECTED_COLUMNS if c not in df.columns]
    if missing:
        print(f"  WARNING: source file is missing expected columns: {missing}")

    a1_path = os.path.join(work_dir, f"A1_EXTRACTION_TMS_OCBA_{date_tag}.csv")
    ordered_cols = [c for c in EXPECTED_COLUMNS if c in df.columns]
    ordered_cols += [c for c in df.columns if c not in ordered_cols]
    df = df[ordered_cols]
    df.to_csv(a1_path, index=False, sep=";")

    return df, a1_path


def step4_copy_to_txt(a1_path, work_dir, date_tag):
    """Spec step 4: copy A1 csv, rename, change extension to .txt -> A2"""
    a2_path = os.path.join(work_dir, f"A2_EXTRACTION_TMS_OCBA_{date_tag}.txt")
    shutil.copyfile(a1_path, a2_path)
    return a2_path


def step5_text_import_wizard(a2_path):
    """
    Spec step 5: "Open this txt file with Excel & Transform it via text
    import wizard: check 'Delimited' > Next > check only 'semicolon' >
    Next > Select column for Serial Numbers (PCL_NO_SERIE_LOT) and
    choose 'Text' > Finish"

    Reproduced by re-reading the semicolon-delimited A2 file with
    PCL_NO_SERIE_LOT (and other identifier columns) forced to text.
    """
    dtype_map = {col: str for col in TEXT_COLUMNS}
    df = pd.read_csv(
        a2_path,
        sep=";",
        dtype=dtype_map,
        encoding="utf-8",  # A2 was written by step3/pandas, so it's UTF-8
        keep_default_na=True,
        quotechar='"',
    )
    return df


def step6_save_as_table_xlsx(df, work_dir, date_tag):
    """
    Spec step 6: "Create a table with headers. Save the file in xlsx
    format as: A3_EXTRACTION_TMS_OCBA_YYMMDD.xlsx"

    Writes a genuine Excel Table object, with the Serial Number
    column's number format set to Text ("@").
    """
    a3_path = os.path.join(work_dir, f"A3_EXTRACTION_TMS_OCBA_{date_tag}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = f"A2_EXTRACTION_TMS_OCBA_{date_tag}"[:31]  # Excel sheet name limit

    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows = len(df) + 1  # +1 for header
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)
    table_ref = f"A1:{last_col_letter}{n_rows}"

    table = Table(displayName="Table1", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    if "PCL_NO_SERIE_LOT" in headers:
        col_idx = headers.index("PCL_NO_SERIE_LOT") + 1  # 1-based
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:  # skip header row
            cell.number_format = "@"

    for i, col_name in enumerate(headers, start=1):
        width = max(12, min(40, len(str(col_name)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width

    wb.save(a3_path)
    return a3_path


def main():
    cfg = load_config()
    paths = get_paths(cfg)

    explicit_path = sys.argv[1] if len(sys.argv) > 1 else None

    if explicit_path:
        source_path = explicit_path
        date_tag = extract_date_tag(explicit_path)
        if date_tag is None:
            raise ValueError(
                f"Could not find a YYMMDD date tag in filename '{explicit_path}'. "
                f"Expected something like EXTRACTION_TMS_OCBA_260623.csv."
            )
        print(f"Using explicitly given file: {source_path} (date tag {date_tag})")
    else:
        source_path, date_tag = find_latest_file_by_date(paths["download_root"])
        print(f"Latest file by date across '{paths['download_root']}': "
              f"{source_path} (date tag {date_tag})")

    # Work folder is named after the ISO week BEFORE the one this
    # script is RUN in (e.g. "Y26W25" if run during week 26), not the
    # batch's own YYMMDD date tag - same convention as
    # 04_validate.py/05_export.py's output folders, so every stage of
    # the pipeline organizes files the same way. The A1/A2/A3
    # filenames still carry the batch's own date tag.
    week_tag = compute_week_tag()
    dated_work_dir = get_dated_subdir(paths["work_dir"], week_tag)
    print(f"Using work folder: {dated_work_dir}  (batch date tag: {date_tag})")

    _, a1_path = step3_save_a1(source_path, dated_work_dir, date_tag)
    print(f"\nStep 3 done. A1 file saved to: {a1_path}")

    a2_path = step4_copy_to_txt(a1_path, dated_work_dir, date_tag)
    print(f"Step 4 done. A2 file (txt copy) saved to: {a2_path}")

    wizard_df = step5_text_import_wizard(a2_path)
    print(f"Step 5 done. Text import wizard equivalent applied "
          f"(semicolon-delimited, PCL_NO_SERIE_LOT forced to Text).")

    a3_path = step6_save_as_table_xlsx(wizard_df, dated_work_dir, date_tag)
    print(f"Step 6 done. A3 table workbook saved to: {a3_path}")

    print(f"\nAll pre-processing steps complete ({len(wizard_df)} rows) for batch {date_tag}.")
    print("\nNext step: run 03_clean_transform.py")


if __name__ == "__main__":
    main()
