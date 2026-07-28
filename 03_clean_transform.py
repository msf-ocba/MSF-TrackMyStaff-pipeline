"""
03_clean_transform.py
-----------------------
Processes the A3 workbook by adding the 'Family' and 'Keep?' columns,
then saves the result as A4.

The script identifies records that should be excluded from later stages
of the import process using three validation checks.

VALIDATION

Each row is evaluated against the following criteria:

  - Family Check
      Derives the Family value from the first four characters of
      FCL_ART_CODE. Articles belonging to the Kit family are marked
      as "NOT ASSET".

  - Customer Check
      Confirms the customer code represents a deployed OC (Operational Center).
      Non-matching records are marked as "NOT DEPLOYED".

  - Mission Check
      Verifies the mission code prefix against the list in
      Mission_code_prefixes.txt. Prefixes not found in the list are
      marked as "NOT DEPLOYED".

KEEP? COLUMN

The Keep? column records any validation results for each row:

  - Blank: the row passed all validation checks.
  - "NOT ASSET": excluded because it is not an asset.
  - "NOT DEPLOYED": excluded because it is not associated with a
    deployed OC (Operational Center) and Mission.
  - Multiple reasons are recorded together where applicable.

OUTPUT

The script creates:

  - A4_EXTRACTION_TMS_OCBA_YYMMDD.xlsx

The workbook contains all original A3 columns, together with the new
Family and Keep? columns, and is saved to the working folder and the
configured output folder.

Usage:
    python 03_clean_transform.py [path_to_A3_workbook.xlsx]

If no A3 workbook is supplied, the script automatically selects the
most recent A3 file.
"""
import os
import sys
import glob
import shutil
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

from common_config import load_config, get_paths, get_dated_subdir, compute_week_tag

REQUIRED_SOURCE_COLUMNS = [
    "FCL_ART_CODE",
    "FCT_CLI_CODE_FAC",
    "PCT_CLI_CODE_LIV",
]

PREFIX_ALLOWLIST_FILE = "./templates/Mission_code_prefixes.txt"  # fallback default; see config.conf [templates] mission_code_prefixes_path

NOT_ASSET = "NOT ASSET"
NOT_DEPLOYED = "NOT DEPLOYED"


def find_latest_a3(work_dir):
    """
    work_dir is the dated-subfolder ROOT (e.g. ./work/). Each batch's
    A3 lives in its own dated subfolder, e.g. ./work/260625/A3_....xlsx.

    We prefer the most recently modified A3 found INSIDE a dated
    subfolder. A flat ./work/A3_....xlsx (not inside any dated folder)
    is only used as a last-resort fallback, and is treated as
    potentially stale - e.g. left over from an old/manual run, or from
    before this script started using dated subfolders. Silently
    preferring a stale flat file over a fresh dated one previously
    caused 03 to read the wrong batch, so dated subfolders always win
    when both exist.
    """
    subfolders = [f for f in glob.glob(os.path.join(work_dir, "*")) if os.path.isdir(f)]
    subfolders.sort(key=os.path.getmtime, reverse=True)
    for folder in subfolders:
        candidates = sorted(glob.glob(os.path.join(folder, "A3_EXTRACTION_TMS_OCBA_*.xlsx")))
        if candidates:
            chosen = max(candidates, key=os.path.getmtime)
            direct = glob.glob(os.path.join(work_dir, "A3_EXTRACTION_TMS_OCBA_*.xlsx"))
            if direct:
                print(f"  NOTE: ignoring {len(direct)} A3 file(s) sitting directly in "
                      f"'{work_dir}' (not inside a dated subfolder) - these look stale "
                      f"and are not used. Using the dated subfolder copy instead: {chosen}")
            return chosen

    # No dated subfolder has an A3 at all - fall back to a flat file, if any.
    direct = sorted(glob.glob(os.path.join(work_dir, "A3_EXTRACTION_TMS_OCBA_*.xlsx")))
    if direct:
        chosen = max(direct, key=os.path.getmtime)
        print(f"  WARNING: no dated subfolder under '{work_dir}' contains an A3 file. "
              f"Falling back to a flat file found directly in '{work_dir}': {chosen}. "
              f"This is likely a leftover from before dated folders were introduced - "
              f"re-run 02_read_files.py to regenerate it properly.")
        return chosen

    raise FileNotFoundError(
        f"No A3_EXTRACTION_TMS_OCBA_*.xlsx file found under '{work_dir}' "
        f"(checked dated subfolders and the flat folder). Run 02_read_files.py first."
    )


def check_required_columns(df):
    missing = [c for c in REQUIRED_SOURCE_COLUMNS if c not in df.columns]
    if missing:
        raise KeyError(
            f"A3 workbook is missing expected source columns: {missing}. "
            f"Check the source export headers match the MSF Logistique field names."
        )


def load_mission_code_prefixes(path):
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Mission code prefix allowlist not found at '{path}'. "
            f"Create it (one 2-character prefix per line) before running this script."
        )
    prefixes = set()
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            prefixes.add(line.upper())
    if not prefixes:
        print(f"  WARNING: '{path}' contains no prefixes - every row will fail the "
              f"mission code check.")
    return prefixes


def compute_family(article_code):
    if pd.isna(article_code):
        return ""
    return str(article_code).strip()[:4]


def compute_keep(row, valid_prefixes, customer_code):
    reasons = []

    family = row["Family"]
    if family.upper().startswith("K"):
        reasons.append(NOT_ASSET)

    fct_cli = str(row.get("FCT_CLI_CODE_FAC", "")).strip()
    if fct_cli != customer_code:
        reasons.append(NOT_DEPLOYED)

    mission_code = str(row.get("PCT_CLI_CODE_LIV", "")).strip()
    mission_prefix = mission_code[:2].upper()
    if mission_prefix not in valid_prefixes:
        reasons.append(NOT_DEPLOYED)

    return "; ".join(reasons)  # blank string if all checks passed


def build_a4_dataframe(df, valid_prefixes, customer_code):
    out = df.copy()
    out["Family"] = out["FCL_ART_CODE"].apply(compute_family)
    out["Keep?"] = out.apply(lambda r: compute_keep(r, valid_prefixes, customer_code), axis=1)

    # Match B's column order: KEEP?, Family, then the original fields.
    other_cols = [c for c in out.columns if c not in ("Family", "Keep?")]
    out = out[["Keep?", "Family"] + other_cols]
    return out


def add_keep_conditional_formatting(ws, n_rows, n_cols, keep_col_idx=1):
    """
    Color every row based on the value in the 'Keep?' column (column A
    by construction - see build_a4_dataframe). Rules are formula-based
    (not one-off cell fills), so they keep working if someone edits
    Keep? by hand later in Excel.

    Priority (most specific first, stop_if_true so only one fires per
    row):
      1. Both "NOT ASSET" and "NOT DEPLOYED" present -> red
      2. "NOT ASSET" only (kit)                       -> purple/lavender
      3. "NOT DEPLOYED" only                          -> yellow
      4. Blank (nothing disqualified it yet - NOT the -> medium gray
         same as a confirmed keep, still needs review)

    Kit rows use a purple/lavender fill instead of orange - orange sat
    too close to the yellow "NOT DEPLOYED" fill to tell apart at a
    glance. The "pending" fill was also darkened from a near-white
    F2F2F2 (practically invisible against the sheet's white background)
    to a clearly visible medium gray.
    """
    keep_col_letter = get_column_letter(keep_col_idx)
    last_col_letter = get_column_letter(n_cols)
    data_range = f"A2:{last_col_letter}{n_rows}"
    ref_cell = f"${keep_col_letter}2"

    fill_both = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_not_asset = PatternFill(start_color="CC99FF", end_color="CC99FF", fill_type="solid")
    fill_not_deployed = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    fill_pending = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")

    rules = [
        FormulaRule(
            formula=[f'AND(ISNUMBER(SEARCH("{NOT_ASSET}",{ref_cell})),'
                     f'ISNUMBER(SEARCH("{NOT_DEPLOYED}",{ref_cell})))'],
            fill=fill_both, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{NOT_ASSET}",{ref_cell}))'],
            fill=fill_not_asset, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{NOT_DEPLOYED}",{ref_cell}))'],
            fill=fill_not_deployed, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'{ref_cell}=""'],
            fill=fill_pending, stopIfTrue=True,
        ),
    ]
    for rule in rules:
        ws.conditional_formatting.add(data_range, rule)


def save_as_table_xlsx(df, work_dir, date_tag):
    a4_path = os.path.join(work_dir, f"A4_EXTRACTION_TMS_OCBA_{date_tag}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = f"A4_EXTRACTION_TMS_OCBA_{date_tag}"[:31]

    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows = len(df) + 1
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)
    table_ref = f"A1:{last_col_letter}{n_rows}"

    table = Table(displayName="Table1", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    if "PCL_NO_SERIE_LOT" in headers:
        col_idx = headers.index("PCL_NO_SERIE_LOT") + 1
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:
            cell.number_format = "@"

    for i, col_name in enumerate(headers, start=1):
        width = max(12, min(40, len(str(col_name)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width

    keep_col_idx = headers.index("Keep?") + 1
    add_keep_conditional_formatting(ws, n_rows, n_cols, keep_col_idx=keep_col_idx)

    try:
        wb.save(a4_path)
    except PermissionError:
        raise PermissionError(
            f"\nCould not save '{a4_path}' - permission denied.\n"
            f"This almost always means a file with that name is currently open "
            f"in Excel - close it and re-run."
        )
    return a4_path


def save_to_output_dir(a4_path, paths):
    """
    Send an identically-named copy of the A4 workbook to the output
    folder, in addition to the dated work-folder copy that
    save_as_table_xlsx() already wrote. It is NOT renamed with a
    "_copy" suffix or similar - same filename, just a second location -
    so downstream steps looking for A4_EXTRACTION_TMS_OCBA_YYMMDD.xlsx
    find it in either place.

    Mirrors the work_dir layout: output goes into a dated subfolder
    (output_dir/<week_tag>/), NEVER flat into output_dir itself. The
    subfolder is named using the SAME shared week_tag convention as
    every other stage of the pipeline (e.g. "Y26W25" - the ISO week
    BEFORE the one this script is run in, per
    common_config.compute_week_tag()) - NOT the batch's own YYMMDD
    date_tag (e.g. "260629") - so this script's output lands in the
    exact same dated folder that 04_validate.py and 05_export.py write
    their own outputs into for the same batch.
    """
    week_tag = compute_week_tag()
    dated_output_dir = get_dated_subdir(paths["output_dir"], week_tag)
    dest_path = os.path.join(dated_output_dir, os.path.basename(a4_path))
    shutil.copy2(a4_path, dest_path)
    return dest_path


def main():
    cfg = load_config()
    paths = get_paths(cfg)

    a3_arg = sys.argv[1] if len(sys.argv) > 1 else None
    a3_path = a3_arg or find_latest_a3(paths["work_dir"])
    print(f"Reading A3 workbook: {a3_path}")

    work_dir_for_output = os.path.dirname(a3_path)  # same dated folder A3 lives in

    # Recover the date tag from the A3 filename so A4 is named consistently.
    base = os.path.basename(a3_path)
    date_tag = base.replace("A3_EXTRACTION_TMS_OCBA_", "").replace(".xlsx", "")

    try:
        df = pd.read_excel(a3_path, dtype=str, keep_default_na=True)
    except PermissionError:
        raise PermissionError(
            f"\nCould not open '{a3_path}' - permission denied.\n"
            f"This almost always means the file is currently open in Excel "
            f"(Windows locks it while open) - close it and try again. If it's "
            f"not open, it may also be mid-sync via OneDrive/SharePoint; wait a "
            f"moment and retry."
        )
    check_required_columns(df)

    customer_code = cfg.get("extraction_rules", "customer_code", fallback="ES001MES")
    prefix_allowlist_path = cfg.get(
        "templates", "mission_code_prefixes_path", fallback=PREFIX_ALLOWLIST_FILE
    )
    valid_prefixes = load_mission_code_prefixes(prefix_allowlist_path)
    print(f"Loaded {len(valid_prefixes)} valid mission code prefix(es): {sorted(valid_prefixes)}")

    a4_df = build_a4_dataframe(df, valid_prefixes, customer_code)

    n_not_asset = a4_df["Keep?"].str.contains(NOT_ASSET, na=False).sum()
    n_not_deployed = a4_df["Keep?"].str.contains(NOT_DEPLOYED, na=False).sum()
    n_blank = (a4_df["Keep?"] == "").sum()
    print(f"Rows: {len(a4_df)} total | {n_blank} passed all checks (blank Keep?) | "
          f"{n_not_asset} flagged NOT ASSET | {n_not_deployed} flagged NOT DEPLOYED "
          f"(a row can be flagged for more than one reason)")

    a4_path = save_as_table_xlsx(a4_df, work_dir_for_output, date_tag)
    print(f"\nA4 workbook saved to: {a4_path}")

    output_path = save_to_output_dir(a4_path, paths)
    print(f"A4 workbook also sent to output folder: {output_path}")

    print("\nNext step: run 04_validate.py")


if __name__ == "__main__":
    main()
