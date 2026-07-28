"""
05_export.py
--------------
Takes this batch's A5 workbook and produces the final import file,
A6_EXTRACTION_TMS_OCBA_YYMMDD.xlsx, mapped to the structure of
C_adminTemplate_IMPORT_asset_MSF_LOG_EN.xlsx for import into TMS.

The script also creates a metadata file,
A6_META_EXTRACTION_TMS_OCBA_YYMMDD.json, alongside A6. This records the
batch summary, warning details and output file paths used by
06_notify.py to generate the completion email.

VALIDATION

Before exporting, the script performs several checks and records any
issues found:

  - Article Code Check: identifies rows still marked for review.
  - Location Check: matches the mission code to the appropriate TMS
    location using templates/location.xlsx. If no match is found, the
    Location field is left blank and the row is flagged.
  - Duplicate Check: identifies duplicate combinations of Article Code
    and Serial Number within the batch.

These checks are informational only and do not prevent A6 from being
generated.

A6 FORMATTING

The exported workbook is formatted to make review easier:

  - Columns are automatically sized to fit their contents, with long
    values wrapped where necessary.
  - Rows are colour-coded according to their status:
      - Green: no issues found.
      - Yellow: Article Code requires review.
      - Orange: missing Location.
      - Red: duplicate Article Code and Serial Number.

OUTPUT

The script creates:

  - A6_EXTRACTION_TMS_OCBA_YYMMDD.xlsx – the final TMS import file.
  - A6_META_EXTRACTION_TMS_OCBA_YYMMDD.json – metadata used by
    06_notify.py when composing the completion email.

COLUMN MAPPING (A5 → A6)

  A6 column              <- A5 source / rule
  ------------------        --------------------------------------------
  Article Code           <- FCL_ART_CODE
  Location               <- Lookup from location.xlsx using the first
                             two characters of PCT_CLI_CODE_LIV
  Status                 <- "In Transit"
  Serial number          <- PCL_NO_SERIE_LOT
  Project Code           <- First five characters of PCT_CLI_CODE_LIV
  Brand                  <- MARQUE
  Model                  <- MODEL (maximum 50 characters)
  Comment                <- Shipment date and Kit reference (when
                             available)
  Purchase reference     <- PCT_REF_CMDE1
  Currency               <- "EUR"
  Price                  <- PRIX_VENTE_UNIT (decimal separator
                             normalised)
  Invoice Number         <- FCT_NO_FACTURE
  Warranty               <- GARANTIE in months
  Manufacturing date     <- DATE_FABRICATION

Columns not required by the TMS import template are omitted from A6.

Usage:
    python 05_export.py [path_to_A5_workbook.xlsx] [path_to_location.xlsx]

If no A5 workbook is specified, the script automatically selects the
most recent A5 file. If no location file is supplied, the path defined
in config.conf is used.

After the export completes, run:

    python 06_notify.py

to send the batch completion email.
"""
import os
import re
import sys
import glob
import json
import math
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Alignment

from common_config import load_config, get_paths, get_dated_subdir, compute_week_tag

LOCATION_LIST_DEFAULT = "./templates/location.xlsx"

CHECKED = "Checked"
REVIEW = "Review"

# C template's column order (Sheet1, row 1 headers) - A6 must match this
# exactly, since this file is what gets imported straight into TMS.
# "Barcode", "Article Description (NOT to be imported)" and
# "Intern Ref or HQ Id" are dropped entirely (see module docstring).
A6_COLUMNS = [
    "Article Code",
    "Location",
    "Status",
    "Serial number",
    "Project Code",
    "Brand",
    "Model",
    "Comment",
    "Purchase reference",
    "Currency",
    "Price",
    "Invoice Number",
    "Warranty",
    "Manufacturing date",
]

# Matches a 2-letter mission code immediately followed by digits inside
# a location path, e.g. ".../CF101-Bangui-.../..." -> "CF". This is how
# templates/location.xlsx's paths encode their prefix - there's no
# separate "prefix" column, the prefix is embedded in the path itself.
LOCATION_PREFIX_RE = re.compile(r"/([A-Za-z]{2})\d+")

# Row-highlight fills for the conditional formatting on A6, using a
# distinctive palette so each state is unmistakable at a glance. If a
# row is flagged by more than one of the three warning checks, the
# most severe fill wins (checked in this order: review, then missing
# location, then duplicate - duplicate is applied last so it takes
# priority when it overlaps with the others). Rows flagged by none of
# the three checks get the green CLEAN_FILL instead.
REVIEW_FILL = PatternFill(start_color="FFEB3B", end_color="FFEB3B", fill_type="solid")           # amber/yellow
MISSING_LOCATION_FILL = PatternFill(start_color="FF9800", end_color="FF9800", fill_type="solid")  # orange
DUPLICATE_FILL = PatternFill(start_color="F44336", end_color="F44336", fill_type="solid")         # red
CLEAN_FILL = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")             # green

MIN_COLUMN_WIDTH = 10
MAX_COLUMN_WIDTH = 50  # cap so one long outlier value doesn't blow out the sheet

# TMS's own import template enforces hard character limits on certain
# fields (confirmed by TMS rejecting the import with "Field 'Model'
# must not be longer than '50' characters"). Any A6 value longer than
# the limit below is truncated before writing, since TMS will reject
# the entire import otherwise. Extend this dict if TMS reports the
# same error for other fields in future.
TMS_FIELD_MAX_LENGTHS = {
    "Model": 50,
}

# Default row height (Excel's default, in points) used as the basis for
# growing a row when it contains a wrapped cell.
DEFAULT_ROW_HEIGHT = 15


def find_latest_a5(work_dir):
    """Same dated-subfolder convention as scripts 3/4's find_latest_aN()."""
    subfolders = [f for f in glob.glob(os.path.join(work_dir, "*")) if os.path.isdir(f)]
    subfolders.sort(key=os.path.getmtime, reverse=True)
    for folder in subfolders:
        candidates = sorted(glob.glob(os.path.join(folder, "A5_EXTRACTION_TMS_OCBA_*.xlsx")))
        if candidates:
            chosen = max(candidates, key=os.path.getmtime)
            direct = glob.glob(os.path.join(work_dir, "A5_EXTRACTION_TMS_OCBA_*.xlsx"))
            if direct:
                print(f"  NOTE: ignoring {len(direct)} A5 file(s) sitting directly in "
                      f"'{work_dir}' (not inside a dated subfolder) - these look stale "
                      f"and are not used. Using the dated subfolder copy instead: {chosen}")
            return chosen

    direct = sorted(glob.glob(os.path.join(work_dir, "A5_EXTRACTION_TMS_OCBA_*.xlsx")))
    if direct:
        chosen = max(direct, key=os.path.getmtime)
        print(f"  WARNING: no dated subfolder under '{work_dir}' contains an A5 file. "
              f"Falling back to a flat file found directly in '{work_dir}': {chosen}. "
              f"This is likely a leftover from before dated folders were introduced - "
              f"re-run 02/03/04 to regenerate it properly.")
        return chosen

    raise FileNotFoundError(
        f"No A5_EXTRACTION_TMS_OCBA_*.xlsx file found under '{work_dir}' "
        f"(checked dated subfolders and the flat folder). Run 04_validate.py first."
    )


def build_location_lookup(location_path):
    """
    Reads templates/location.xlsx (a single column of full location
    paths) and builds a {2-letter prefix: location path} dict by
    extracting the mission-code segment from each path, e.g.:

        "OCBA/Assets/CAR/CF101-Bangui-Coordination/VL_Transit_IN"
          -> prefix "CF"

    If the same prefix appears more than once with DIFFERENT paths,
    that's ambiguous - we keep the first and warn, since silently
    picking one could send an asset to the wrong location.
    """
    if not os.path.exists(location_path):
        raise FileNotFoundError(
            f"Location list not found at '{location_path}'. Put MSF Logistique's "
            f"transit-in location list there (one full location path per row), "
            f"or pass a path as the 2nd argument."
        )

    df = pd.read_excel(location_path, dtype=str, header=0)
    first_col = df.columns[0]
    paths = df[first_col].dropna().astype(str).str.strip()
    paths = paths[paths != ""]
    # Drop any stray zero-width-space-only rows seen in real exports
    paths = paths[~paths.str.fullmatch(r"[\u200b\s]*")]

    lookup = {}
    ambiguous = []
    for path in paths:
        match = LOCATION_PREFIX_RE.search(path)
        if not match:
            print(f"  WARNING: could not extract a 2-letter prefix from location "
                  f"path '{path}' - skipping this entry.")
            continue
        prefix = match.group(1).upper()
        if prefix in lookup and lookup[prefix] != path:
            ambiguous.append(prefix)
        lookup.setdefault(prefix, path)

    if ambiguous:
        print(f"  WARNING: these prefixes appear more than once with DIFFERENT "
              f"location paths in '{location_path}': {sorted(set(ambiguous))}. "
              f"Using the first one found for each - check the file for duplicates.")

    return lookup


def parse_eu_price(value):
    """Convert '491,96' (EU comma-decimal) -> '491.96' (string, as the
    template's own SUBSTITUTE formula does - kept as text, not float,
    since that's what the live formula produces)."""
    if pd.isna(value):
        return ""
    return str(value).strip().replace(",", ".")


def build_comment(row):
    shipment_date = row.get("SDT_DT_ENLEV")
    date_str = ""
    if pd.notna(shipment_date) and str(shipment_date).strip():
        # SDT_DT_ENLEV arrives as DD/MM/YYYY text already (see 02_read_files.py)
        date_str = str(shipment_date).strip()

    comment = f"Shipped from BDO on {date_str}"

    kit_code = row.get("CODE_KIT")
    if pd.notna(kit_code) and str(kit_code).strip():
        comment += f"\nwithin Kit {str(kit_code).strip()}"

    return comment


def build_a6_dataframe(df, location_lookup):
    """
    Builds the A6 dataframe and returns (a6_df, missing_location_mask,
    duplicate_mask) - all three share df's original index, so they
    line up row-for-row with the source A5 dataframe.

    duplicate_mask flags a row only when ANOTHER row in the batch has
    BOTH the same Article Code AND the same Serial number (case/
    whitespace insensitive on both). Article Code alone repeating is
    common and not a problem; it's the (Article Code, Serial number)
    PAIR repeating that means TMS will actually reject the import as a
    true duplicate.

    Any field listed in TMS_FIELD_MAX_LENGTHS (e.g. Model) is silently
    truncated to its max length before this returns, since TMS rejects
    the whole import otherwise.
    """
    out = pd.DataFrame(index=df.index)

    out["Article Code"] = df["FCL_ART_CODE"]

    prefixes = df["PCT_CLI_CODE_LIV"].astype(str).str.strip().str.slice(0, 2).str.upper()
    out["Location"] = prefixes.map(location_lookup).fillna("")
    missing_location_mask = (out["Location"] == "") & df["PCT_CLI_CODE_LIV"].notna()

    out["Status"] = "In Transit"
    out["Serial number"] = df["PCL_NO_SERIE_LOT"]
    out["Project Code"] = df["PCT_CLI_CODE_LIV"].astype(str).str.strip().str.slice(0, 5)

    out["Brand"] = df.get("MARQUE", "").fillna("")
    out["Model"] = df.get("MODEL", "").fillna("")

    out["Comment"] = df.apply(build_comment, axis=1)

    out["Purchase reference"] = df["PCT_REF_CMDE1"]
    out["Currency"] = "EUR"
    out["Price"] = df["PRIX_VENTE_UNIT"].apply(parse_eu_price)
    out["Invoice Number"] = df["FCT_NO_FACTURE"]

    garantie = df.get("GARANTIE")
    out["Warranty"] = garantie.apply(
        lambda v: f"{str(v).strip()} months" if pd.notna(v) and str(v).strip() != "" else ""
    ) if garantie is not None else ""

    out["Manufacturing date"] = df.get("DATE_FABRICATION", "").fillna("")

    out = out[A6_COLUMNS]

    # TMS field-length enforcement: truncate any field TMS is known to
    # reject past a certain length (see TMS_FIELD_MAX_LENGTHS), so the
    # import doesn't fail outright.
    for col_name, max_len in TMS_FIELD_MAX_LENGTHS.items():
        if col_name not in out.columns:
            continue
        out[col_name] = out[col_name].fillna("").astype(str).str.slice(0, max_len)

    # Duplicate Article Code + Serial Number check - case/whitespace
    # insensitive on both fields. A row only counts as a duplicate if
    # some OTHER row shares both its Article Code AND its Serial
    # number; blank Article Codes never count as a "duplicate" of each
    # other, and matching Article Code alone (different Serial
    # numbers) does NOT flag a row.
    article_code_key = out["Article Code"].fillna("").astype(str).str.strip().str.upper()
    serial_key = out["Serial number"].fillna("").astype(str).str.strip().str.upper()
    pair_key = article_code_key + "\u241f" + serial_key  # unlikely-to-collide separator
    non_blank = article_code_key != ""
    duplicate_mask = pair_key.duplicated(keep=False) & non_blank

    return out, missing_location_mask, duplicate_mask


def compute_column_widths(values_by_column, headers):
    """
    Content-based column widths so every field can be read without
    manually resizing columns - width follows whichever is longer, the
    header or the longest value in that column - capped at
    MAX_COLUMN_WIDTH so one outlier value doesn't blow out the sheet.
    Cells longer than the cap get text-wrapping applied instead (see
    apply_wrap_text_for_long_values()), so nothing is hidden.
    """
    widths = []
    for col_name in headers:
        header_len = max(len(line) for line in str(col_name).split("\n"))
        col_values = values_by_column.get(col_name, [])
        data_len = max((len(str(v)) for v in col_values), default=0)
        width = max(header_len, data_len) + 2
        width = max(MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, width))
        widths.append(width)
    return widths


def apply_wrap_text_for_long_values(ws, df, headers, column_widths):
    """
    Any cell whose value is longer than MAX_COLUMN_WIDTH characters
    would otherwise be visually clipped by the column-width cap, even
    though the value is still fully present in the file (readable via
    the formula bar). Instead of leaving it clipped, turn on text
    wrapping for that specific cell so the full value is visible
    directly in the sheet, and grow that row's height to fit the
    number of wrapped lines it now needs.

    This ONLY changes how long values are displayed - it does not
    alter, shorten, or touch the underlying value written to the cell,
    since A6 is imported straight into TMS and its data must stay
    exactly as computed.
    """
    row_heights = {}  # excel row number -> max lines needed across its cells

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        col_width = column_widths[col_idx - 1]
        values = df[col_name].tolist()
        for row_offset, value in enumerate(values):
            text = "" if value is None else str(value)
            if not text:
                continue

            # A value already containing newlines (e.g. the multi-line
            # Comment field) wraps at those newlines regardless of
            # length, so it also benefits from wrap_text once any one
            # of its lines exceeds the column width.
            longest_line = max(len(line) for line in text.split("\n"))
            if longest_line <= MAX_COLUMN_WIDTH and "\n" not in text:
                continue

            excel_row = row_offset + 2  # +1 for header, +1 for 1-indexing
            cell = ws.cell(row=excel_row, column=col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            # Estimate wrapped line count: existing newlines each force
            # a break, plus each of those segments may itself wrap
            # across ceil(len/col_width) lines.
            lines_needed = 0
            for line in text.split("\n"):
                lines_needed += max(1, math.ceil(len(line) / col_width))
            row_heights[excel_row] = max(row_heights.get(excel_row, 1), lines_needed)

    for excel_row, lines_needed in row_heights.items():
        ws.row_dimensions[excel_row].height = DEFAULT_ROW_HEIGHT * lines_needed


def save_as_table_xlsx(df, output_dir, date_tag, review_mask, missing_location_mask, duplicate_mask):
    """
    Writes A6 as an Excel table, with:
      - content-based column autofit (capped at MAX_COLUMN_WIDTH), with
        wrap-text applied to any individual cell that exceeds the cap
        so long values stay fully visible instead of being clipped,
      - entire-row highlight fills for every row: the three warning
        checks (review / missing location / duplicate Article Code +
        Serial number) each get their own distinctive colour, and any
        row flagged by none of them gets a green "all good" highlight.
    review_mask, missing_location_mask and duplicate_mask must share
    df's index/order.
    """
    a6_path = os.path.join(output_dir, f"A6_EXTRACTION_TMS_OCBA_{date_tag}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows = len(df) + 1
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)
    table_ref = f"A1:{last_col_letter}{n_rows}"

    table = Table(displayName="Table1", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    serial_col_name = "Serial number"
    if serial_col_name in headers:
        col_idx = headers.index(serial_col_name) + 1
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:
            cell.number_format = "@"

    # Content-based, capped-at-50 column autofit.
    values_by_column = {col: df[col].tolist() for col in headers}
    widths = compute_column_widths(values_by_column, headers)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Any value that got clipped by the width cap above gets wrap-text
    # turned on instead, so it's still fully readable in the sheet.
    apply_wrap_text_for_long_values(ws, df, headers, widths)

    # Entire-row conditional highlighting. Positional arrays (reset_index)
    # so they line up with the itertuples() write order above regardless
    # of the original dataframe's index labels. Every row gets exactly
    # one fill: the most severe warning colour if flagged, otherwise the
    # green "clean" colour.
    review_arr = review_mask.reset_index(drop=True)
    missing_loc_arr = missing_location_mask.reset_index(drop=True)
    duplicate_arr = duplicate_mask.reset_index(drop=True)

    for i in range(len(df)):
        fill = CLEAN_FILL
        if bool(review_arr.iloc[i]):
            fill = REVIEW_FILL
        if bool(missing_loc_arr.iloc[i]):
            fill = MISSING_LOCATION_FILL
        if bool(duplicate_arr.iloc[i]):
            fill = DUPLICATE_FILL
        excel_row = i + 2  # +1 for header, +1 for 1-indexing
        for col in range(1, n_cols + 1):
            ws.cell(row=excel_row, column=col).fill = fill

    try:
        wb.save(a6_path)
    except PermissionError:
        raise PermissionError(
            f"\nCould not save '{a6_path}' - permission denied.\n"
            f"This almost always means a file with that name is currently open "
            f"in Excel - close it and re-run."
        )
    return a6_path


def write_meta_json(output_dir, date_tag, meta):
    """
    Writes A6_META_EXTRACTION_TMS_OCBA_<date_tag>.json into the dated
    output folder - the handoff file 06_notify.py reads to compose and
    send the completion email, so it doesn't need to re-read A5 or
    re-run any of the checks above.
    """
    meta_path = os.path.join(output_dir, f"A6_META_EXTRACTION_TMS_OCBA_{date_tag}.json")
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)
    return meta_path


def main():
    cfg = load_config()
    paths = get_paths(cfg)

    a5_arg = sys.argv[1] if len(sys.argv) > 1 else None
    location_arg = sys.argv[2] if len(sys.argv) > 2 else None

    a5_path = a5_arg or find_latest_a5(paths["work_dir"])
    print(f"Reading A5 workbook: {a5_path}")

    location_path = location_arg or cfg.get(
        "templates", "location_list_path", fallback=LOCATION_LIST_DEFAULT
    )

    base = os.path.basename(a5_path)
    date_tag = base.replace("A5_EXTRACTION_TMS_OCBA_", "").replace(".xlsx", "")

    try:
        df = pd.read_excel(a5_path, dtype=str, keep_default_na=True)
    except PermissionError:
        raise PermissionError(
            f"\nCould not open '{a5_path}' - permission denied.\n"
            f"This almost always means the file is currently open in Excel "
            f"(Windows locks it while open) - close it and try again. If it's "
            f"not open, it may also be mid-sync via OneDrive/SharePoint; wait a "
            f"moment and retry."
        )

    required = ["FCL_ART_CODE", "PCL_NO_SERIE_LOT", "PCT_CLI_CODE_LIV", "Article Code Check"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise KeyError(f"A5 workbook is missing expected column(s): {missing}.")

    print(f"A5 rows: {len(df)}")

    review_mask = df["Article Code Check"].astype(str).str.strip() == REVIEW
    n_review = int(review_mask.sum())
    review_rows_summary = [
        f"Row with Article Code '{code}' (Serial '{serial}') needs manual TMS lookup"
        for code, serial in zip(df.loc[review_mask, "FCL_ART_CODE"], df.loc[review_mask, "PCL_NO_SERIE_LOT"])
    ]
    if n_review:
        print(f"  WARNING: {n_review} row(s) still have Article Code Check = 'Review'. "
              f"They are INCLUDED in A6 anyway (current process does not block export "
              f"on this), but should be fixed in TMS/A5 as soon as possible.")

    location_lookup = build_location_lookup(location_path)
    print(f"Loaded {len(location_lookup)} location prefix(es) from: {location_path}")

    a6_df, missing_location_mask, duplicate_mask = build_a6_dataframe(df, location_lookup)

    n_missing_location = int(missing_location_mask.sum())
    missing_location_summary = [
        f"Row with Article Code '{code}' (Project Code '{proj}') - no location match "
        f"for prefix '{str(proj).strip()[:2].upper()}'"
        for code, proj in zip(
            df.loc[missing_location_mask, "FCL_ART_CODE"],
            df.loc[missing_location_mask, "PCT_CLI_CODE_LIV"],
        )
    ]
    if n_missing_location:
        print(f"  WARNING: {n_missing_location} row(s) have no matching Location for "
              f"their Project Code prefix. Location left BLANK for those rows in A6 - "
              f"fix templates/location.xlsx and re-run, or fill them in by hand.")

    n_duplicate = int(duplicate_mask.sum())
    duplicate_summary = [
        f"Row with Article Code '{code}' (Serial '{serial}') - both Article Code AND "
        f"Serial number appear more than once together in this batch"
        for code, serial in zip(
            df.loc[duplicate_mask, "FCL_ART_CODE"],
            df.loc[duplicate_mask, "PCL_NO_SERIE_LOT"],
        )
    ]
    if n_duplicate:
        print(f"  WARNING: {n_duplicate} row(s) share BOTH an Article Code AND a Serial "
              f"number with another row in this batch - TMS will block these as true "
              f"duplicates on import. Not blocked here, but highlighted in A6 and listed "
              f"in the completion email - double-check these aren't duplicate entries of "
              f"the same item.")

    # Output folder is named after the ISO week BEFORE the one this
    # script is run in (e.g. "Y26W25" if run during week 26), not the
    # batch's own date tag - see common_config.compute_week_tag().
    week_tag = compute_week_tag()
    dated_output_dir = get_dated_subdir(paths["output_dir"], week_tag)

    a6_path = save_as_table_xlsx(
        a6_df, dated_output_dir, date_tag,
        review_mask=review_mask,
        missing_location_mask=missing_location_mask,
        duplicate_mask=duplicate_mask,
    )
    print(f"\nA6 (final TMS import file) saved to: {a6_path}")

    # A row can be flagged by more than one check at once, so "clean"
    # rows are those flagged by NONE of them - not just total minus the
    # sum of the three.
    flagged_mask = review_mask | missing_location_mask | duplicate_mask
    n_flagged = int(flagged_mask.sum())
    n_clean = len(df) - n_flagged

    if n_flagged:
        print(f"\n{n_clean} of {len(df)} item(s) are clean; {n_flagged} need review.")
    else:
        print("\nNo review needed - every row passed the Article Code Check, has a "
              "matching Location, and has a unique Article Code + Serial number pair.")

    # Write the handoff file for 06_notify.py, which sends (or prints)
    # the completion email using exactly these numbers/details - no
    # need to re-derive anything from A5 there.
    meta = {
        "week_tag": week_tag,
        "date_tag": date_tag,
        "output_dir": dated_output_dir,
        "total_rows": len(df),
        "n_clean": n_clean,
        "n_review": n_review,
        "n_missing_location": n_missing_location,
        "n_duplicate": n_duplicate,
        "review_rows_summary": review_rows_summary,
        "missing_location_summary": missing_location_summary,
        "duplicate_summary": duplicate_summary,
        "a6_path": a6_path,
    }
    meta_path = write_meta_json(dated_output_dir, date_tag, meta)
    print(f"Meta file for 06_notify.py saved to: {meta_path}")

    print("\nNext step: review any WARNING lines above, then run "
          "'python 06_notify.py' to send the completion email, then "
          "import A6 into TMS.")


if __name__ == "__main__":
    main()
