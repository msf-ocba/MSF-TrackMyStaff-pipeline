"""
04_validate.py
----------------
Processes this batch's A4 workbook by selecting rows where the
'Keep?' column is blank and validating each row's FCL_ART_CODE
against the TMS article list.

ARTICLE CODE VALIDATION

Each article code is checked against the reference data in
templates/TMS_UniDataArticles.xlsx:

  - Article code found -> "Checked"
  - Article code not found -> "Review"

ARTICLE COMPOSE LOOKUP

Some article codes require conversion before they can be matched
against the TMS article list. The script uses
templates/ArticleCompose.xlsx to substitute these codes with their
corresponding TMS article codes before performing validation.

OUTPUT

The script creates up to three files:

  - A5_EXTRACTION_TMS_OCBA_YYMMDD.xlsx
      Contains all rows where 'Keep?' is blank, together with the
      "Article Code Check" result. Rows are colour-coded to highlight
      validation status:
          - Green: Checked
          - Red: Review
      NOT created if there are zero rows with blank 'Keep?' - see
      EXIT CODES below.

  - A5B_EXTRACTION_TMS_OCBA_YYMMDD.xlsx
      Contains all rows excluded from A5 (those where 'Keep?' is not
      blank). The workbook retains the same row colouring applied in
      A4 to indicate why each record was excluded. Always created
      when there is at least one excluded row, even if A5 is not
      created.

  - A5_META_EXTRACTION_TMS_OCBA_YYMMDD.json
      Only written when there are zero rows with blank 'Keep?' (see
      exit code 3 below). This is a lighter counterpart to
      05_export.py's A6_META_....json: since no A6 exists for this
      batch, 06_notify.py reads this file instead so it can still
      send a completion email explaining nothing was imported this
      run - technicians rely on that email arriving, not on checking
      server logs.

EXIT CODES
  0  - Normal. A5 was produced (there was at least one row with
       blank 'Keep?' to validate/import).
  3  - Normal, NOT an error. Every row in this batch's A4 workbook
       had a non-blank 'Keep?' (e.g. all "NOT DEPLOYED" / "NOT
       ASSET"), so there is nothing new to import into TMS this run.
       A5B and A5_META_....json are still written (for the audit
       trail and for 06_notify.py, respectively), but A5 is not
       created and no article validation is performed. A caller (see
       run_pipeline.py) should skip 05_export.py (nothing to export)
       but still run 06_notify.py, so a "nothing to import" email
       still goes out.
  1 (implicit, via uncaught exception) - a real failure.

Usage:
    python 04_validate.py [path_to_A4_workbook.xlsx]

If no A4 workbook is supplied, the script automatically selects the
most recent A4 file.
"""
import os
import sys
import glob
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

from common_config import load_config, get_paths, get_dated_subdir, compute_week_tag

TMS_ARTICLE_LIST_DEFAULT = "./templates/TMS_UniDataArticles.xlsx"
TMS_ARTICLE_CODE_COLUMN = "Article Code"

ARTICLE_COMPOSE_DEFAULT = "./templates/ArticleCompose.xlsx"
ARTICLE_COMPOSE_SOURCE_COL = "Article"
ARTICLE_COMPOSE_TARGET_COL = "ArticleCompose"

CHECKED = "Checked"
REVIEW = "Review"

NOT_ASSET = "NOT ASSET"
NOT_DEPLOYED = "NOT DEPLOYED"

# Exit codes - see module docstring's EXIT CODES section.
EXIT_OK = 0
EXIT_NOTHING_TO_IMPORT = 3


def find_latest_a4(work_dir):
    subfolders = [f for f in glob.glob(os.path.join(work_dir, "*")) if os.path.isdir(f)]
    subfolders.sort(key=os.path.getmtime, reverse=True)
    for folder in subfolders:
        candidates = sorted(glob.glob(os.path.join(folder, "A4_EXTRACTION_TMS_OCBA_*.xlsx")))
        if candidates:
            chosen = max(candidates, key=os.path.getmtime)
            direct = glob.glob(os.path.join(work_dir, "A4_EXTRACTION_TMS_OCBA_*.xlsx"))
            if direct:
                print(f"  NOTE: ignoring {len(direct)} A4 file(s) sitting directly in "
                      f"'{work_dir}' (not inside a dated subfolder) - these look stale "
                      f"and are not used. Using the dated subfolder copy instead: {chosen}")
            return chosen

    direct = sorted(glob.glob(os.path.join(work_dir, "A4_EXTRACTION_TMS_OCBA_*.xlsx")))
    if direct:
        chosen = max(direct, key=os.path.getmtime)
        print(f"  WARNING: no dated subfolder under '{work_dir}' contains an A4 file. "
              f"Falling back to a flat file found directly in '{work_dir}': {chosen}. "
              f"This is likely a leftover from before dated folders were introduced - "
              f"re-run 02/03 to regenerate it properly.")
        return chosen

    raise FileNotFoundError(
        f"No A4_EXTRACTION_TMS_OCBA_*.xlsx file found under '{work_dir}' "
        f"(checked dated subfolders and the flat folder). Run 03_clean_transform.py first."
    )


def resolve_template_path(path):
    if os.path.exists(path):
        return path

    folder = os.path.dirname(path) or "."
    target_name = os.path.basename(path).lower()
    if os.path.isdir(folder):
        for name in os.listdir(folder):
            if name.lower() == target_name:
                resolved = os.path.join(folder, name)
                print(f"  NOTE: found '{name}' (case-insensitive match for "
                      f"'{os.path.basename(path)}') - using that.")
                return resolved

    return path  # nothing found; return original so the caller's error message is accurate


def load_tms_article_codes(path):
    path = resolve_template_path(path)
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"TMS article list not found at '{path}'. Export the current article "
            f"list from TMS and save it there (or pass a path as the 2nd argument), "
            f"keeping the '{TMS_ARTICLE_CODE_COLUMN}' column name."
        )
    df = pd.read_excel(path, dtype=str)
    if TMS_ARTICLE_CODE_COLUMN not in df.columns:
        raise KeyError(
            f"'{path}' has no '{TMS_ARTICLE_CODE_COLUMN}' column. Found columns: "
            f"{list(df.columns)}. Check the export still uses that header name."
        )
    codes = set(df[TMS_ARTICLE_CODE_COLUMN].dropna().astype(str).str.strip())
    return codes


def load_article_compose_map(path):
    path = resolve_template_path(path)
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Article compose template not found at '{path}'. This file maps raw "
            f"extraction codes for 'hidden' Kit articles (no leading 'K') to the "
            f"code TMS actually uses for them. Save it there (or pass a path as "
            f"the 3rd argument), keeping the '{ARTICLE_COMPOSE_SOURCE_COL}' / "
            f"'{ARTICLE_COMPOSE_TARGET_COL}' column names."
        )
    df = pd.read_excel(path, dtype=str)
    missing = [c for c in (ARTICLE_COMPOSE_SOURCE_COL, ARTICLE_COMPOSE_TARGET_COL) if c not in df.columns]
    if missing:
        raise KeyError(
            f"'{path}' is missing expected column(s): {missing}. Found columns: "
            f"{list(df.columns)}."
        )
    df = df.dropna(subset=[ARTICLE_COMPOSE_SOURCE_COL, ARTICLE_COMPOSE_TARGET_COL])
    mapping = {
        str(src).strip(): str(dst).strip()
        for src, dst in zip(df[ARTICLE_COMPOSE_SOURCE_COL], df[ARTICLE_COMPOSE_TARGET_COL])
    }
    return mapping


def apply_article_compose(blank_df, compose_map):
    raw_codes = blank_df["FCL_ART_CODE"].astype(str).str.strip()
    composed = raw_codes.map(compose_map)
    matched_mask = composed.notna()

    n_matched = int(matched_mask.sum())
    if n_matched:
        blank_df.loc[matched_mask, "FCL_ART_CODE"] = composed[matched_mask]

    return blank_df, n_matched


def add_article_check_conditional_formatting(ws, n_rows, n_cols, check_col_idx):
    if n_rows <= 1:
        # No data rows (header only) - there is nothing to format, and
        # an "A2:...1" range is invalid and raises inside openpyxl.
        return

    check_col_letter = get_column_letter(check_col_idx)
    last_col_letter = get_column_letter(n_cols)
    data_range = f"A2:{last_col_letter}{n_rows}"
    ref_cell = f"${check_col_letter}2"

    fill_checked = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_review = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    ws.conditional_formatting.add(
        data_range,
        FormulaRule(formula=[f'{ref_cell}="{CHECKED}"'], fill=fill_checked, stopIfTrue=True),
    )
    ws.conditional_formatting.add(
        data_range,
        FormulaRule(formula=[f'{ref_cell}="{REVIEW}"'], fill=fill_review, stopIfTrue=True),
    )


def add_keep_conditional_formatting(ws, n_rows, n_cols, keep_col_idx):
    """
    Same rules as 03_clean_transform.py's add_keep_conditional_formatting,
    reproduced here so A5B (the excluded-rows workbook) keeps the exact
    same Keep?-based row coloring that A4 has. See that script for the
    rule rationale.

    Kit rows use a purple/lavender fill (CC99FF) instead of orange -
    orange sat too close to the yellow "NOT DEPLOYED" fill to tell
    apart at a glance. The "pending" fill is a medium gray (BFBFBF)
    rather than a near-white, which was practically invisible against
    the sheet's white background. Kept in sync manually with
    03_clean_transform.py.
    """
    if n_rows <= 1:
        # No data rows (header only) - nothing to format.
        return

    keep_col_letter = get_column_letter(keep_col_idx)
    last_col_letter = get_column_letter(n_cols)
    data_range = f"A2:{last_col_letter}{n_rows}"
    ref_cell = f"${keep_col_letter}2"

    fill_both = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_not_asset = PatternFill(start_color="CC99FF", end_color="CC99FF", fill_type="solid")
    fill_not_deployed = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    fill_pending = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")

    rules = [
        FormulaRule(
            formula=[f'AND(ISNUMBER(SEARCH("{NOT_ASSET}",{ref_cell})),'
                     f'ISNUMBER(SEARCH("{NOT_DEPLOYED}",{ref_cell})))'],
            fill=fill_both, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{NOT_ASSET}",{ref_cell}))'],
            fill=fill_not_asset, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{NOT_DEPLOYED}",{ref_cell}))'],
            fill=fill_not_deployed, stopIfTrue=True,
        ),
        FormulaRule(
            formula=[f'{ref_cell}=""'],
            fill=fill_pending, stopIfTrue=True,
        ),
    ]
    for rule in rules:
        ws.conditional_formatting.add(data_range, rule)


def _write_table_workbook(df, out_path, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    n_rows = len(df) + 1
    n_cols = len(headers)
    last_col_letter = get_column_letter(n_cols)
    table_ref = f"A1:{last_col_letter}{n_rows}"

    table = Table(displayName=f"Table_{sheet_title[:20].replace(' ', '_')}", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    if "PCL_NO_SERIE_LOT" in headers:
        col_idx = headers.index("PCL_NO_SERIE_LOT") + 1
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter][1:]:
            cell.number_format = "@"

    for i, col_name in enumerate(headers, start=1):
        width = max(12, min(40, len(str(col_name)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width

    try:
        wb.save(out_path)
    except PermissionError:
        raise PermissionError(
            f"\nCould not save '{out_path}' - permission denied.\n"
            f"This almost always means a file with that name is currently open "
            f"in Excel - close it and re-run."
        )
    return wb, ws, headers, n_rows, n_cols


def save_a5_xlsx(df, work_dir, date_tag):
    a5_path = os.path.join(work_dir, f"A5_EXTRACTION_TMS_OCBA_{date_tag}.xlsx")
    wb, ws, headers, n_rows, n_cols = _write_table_workbook(
        df, a5_path, f"A5_EXTRACTION_TMS_OCBA_{date_tag}"
    )

    check_col_idx = headers.index("Article Code Check") + 1
    add_article_check_conditional_formatting(ws, n_rows, n_cols, check_col_idx)

    wb.save(a5_path)  # re-save with conditional formatting added
    return a5_path


def save_a5b_xlsx(df, output_dir, date_tag):
    a5b_path = os.path.join(output_dir, f"A5B_EXTRACTION_TMS_OCBA_{date_tag}.xlsx")
    wb, ws, headers, n_rows, n_cols = _write_table_workbook(
        df, a5b_path, f"A5B_EXTRACTION_TMS_OCBA_{date_tag}"
    )

    keep_col_idx = headers.index("Keep?") + 1
    add_keep_conditional_formatting(ws, n_rows, n_cols, keep_col_idx)

    wb.save(a5b_path)  # re-save with conditional formatting added
    return a5b_path


def resolve_output_dir(paths):
    week_tag = compute_week_tag()
    output_dir = get_dated_subdir(paths["output_dir"], week_tag)
    return output_dir, week_tag


def write_nothing_to_import_meta(output_dir, date_tag, week_tag, total_rows, a5b_path, a4_path):
    """
    Written only when there are zero rows with blank 'Keep?' (nothing
    to import this run - see EXIT CODES above). This is a lighter
    counterpart to 05_export.py's A6_META_EXTRACTION_TMS_OCBA_*.json:
    there is no A6 to describe, so this instead gives 06_notify.py
    just enough (an "outcome" marker plus the batch numbers and the
    A4/A5B paths) to send a distinct "nothing to import" completion
    email - still attaching A4/A5B - instead of the usual A6-summary
    email. Users depend on getting that email either way, not on
    checking server logs to find out the batch ran.
    """
    meta = {
        "outcome": "nothing_to_import",
        "week_tag": week_tag,
        "date_tag": date_tag,
        "output_dir": output_dir,
        "total_rows": total_rows,
        "a5b_path": a5b_path,
        "a4_path": a4_path,
    }
    meta_path = os.path.join(output_dir, f"A5_META_EXTRACTION_TMS_OCBA_{date_tag}.json")
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)
    return meta_path


def main():
    cfg = load_config()
    paths = get_paths(cfg)

    a4_arg = sys.argv[1] if len(sys.argv) > 1 else None
    tms_arg = sys.argv[2] if len(sys.argv) > 2 else None
    compose_arg = sys.argv[3] if len(sys.argv) > 3 else None

    a4_path = a4_arg or find_latest_a4(paths["work_dir"])
    print(f"Reading A4 workbook: {a4_path}")

    tms_list_path = tms_arg or cfg.get(
        "templates", "tms_article_list_path", fallback=TMS_ARTICLE_LIST_DEFAULT
    )
    article_compose_path = compose_arg or cfg.get(
        "templates", "article_compose_path", fallback=ARTICLE_COMPOSE_DEFAULT
    )

    work_dir_for_output = os.path.dirname(a4_path)
    base = os.path.basename(a4_path)
    date_tag = base.replace("A4_EXTRACTION_TMS_OCBA_", "").replace(".xlsx", "")

    try:
        df = pd.read_excel(a4_path, dtype=str, keep_default_na=True)
    except PermissionError:
        raise PermissionError(
            f"\nCould not open '{a4_path}' - permission denied.\n"
            f"This almost always means the file is currently open in Excel "
            f"(Windows locks it while open) - close it and try again. If it's "
            f"not open, it may also be mid-sync via OneDrive/SharePoint; wait a "
            f"moment and retry."
        )

    if "Keep?" not in df.columns:
        raise KeyError("A4 workbook has no 'Keep?' column - check it was produced by 03_clean_transform.py.")
    if "FCL_ART_CODE" not in df.columns:
        raise KeyError("A4 workbook has no 'FCL_ART_CODE' column.")

    keep_blank_mask = df["Keep?"].isna() | (df["Keep?"].astype(str).str.strip() == "")
    blank_df = df[keep_blank_mask].copy()
    excluded_df = df[~keep_blank_mask].copy()
    print(f"A4 rows: {len(df)} total | {len(blank_df)} with blank Keep? carried forward to A5 | "
          f"{len(excluded_df)} excluded (non-blank Keep?) carried to A5B")

    # --- Nothing to import this run? ---
    # Every row already has a non-blank 'Keep?' (e.g. "NOT DEPLOYED" /
    # "NOT ASSET"), so there is nothing new for TMS. Stop here rather
    # than continuing on to the article-compose/TMS lookups (nothing to
    # validate) and attempting to build an empty A5 workbook (which
    # openpyxl's conditional formatting can't represent - an "A2:...1"
    # range with 0 data rows is invalid and raises).
    if len(blank_df) == 0:
        print("\nNo rows with a blank 'Keep?' in this A4 workbook - there is nothing "
              "new to validate or import into TMS this run.")

        output_dir, week_tag = resolve_output_dir(paths)
        a5b_path = save_a5b_xlsx(excluded_df, output_dir, date_tag)
        print(f"A5B workbook (excluded rows, for the audit trail) saved to: {a5b_path} "
              f" (output folder: {week_tag})")

        meta_path = write_nothing_to_import_meta(
            output_dir, date_tag, week_tag, len(df), a5b_path, a4_path
        )
        print(f"Meta file for 06_notify.py saved to: {meta_path}")

        print("\nNo A5 workbook was produced - there is nothing to import into TMS "
              "this run. 05_export.py will not run, but run_pipeline.py will still "
              "run 06_notify.py so a 'nothing to import' completion email goes out.")
        sys.exit(EXIT_NOTHING_TO_IMPORT)

    # --- Article compose substitution (hidden Kits without leading 'K') ---
    compose_map = load_article_compose_map(article_compose_path)
    print(f"Loaded {len(compose_map)} article-compose mapping(s) from: {article_compose_path}")
    blank_df, n_composed = apply_article_compose(blank_df, compose_map)
    if n_composed:
        print(f"  Substituted {n_composed} row(s)' FCL_ART_CODE with its ArticleCompose "
              f"value before the TMS check.")

    # --- TMS article code check ---
    tms_codes = load_tms_article_codes(tms_list_path)
    print(f"Loaded {len(tms_codes)} known article code(s) from: {tms_list_path}")

    article_codes = blank_df["FCL_ART_CODE"].astype(str).str.strip()
    blank_df["Article Code Check"] = article_codes.apply(
        lambda code: CHECKED if code in tms_codes else REVIEW
    )

    n_checked = (blank_df["Article Code Check"] == CHECKED).sum()
    n_review = (blank_df["Article Code Check"] == REVIEW).sum()
    print(f"Article Code Check: {n_checked} Checked | {n_review} Review (not found in TMS list)")

    a5_path = save_a5_xlsx(blank_df, work_dir_for_output, date_tag)
    print(f"\nA5 workbook saved to: {a5_path}")

    output_dir, week_tag = resolve_output_dir(paths)
    a5b_path = save_a5b_xlsx(excluded_df, output_dir, date_tag)
    print(f"A5B workbook (excluded rows) saved to: {a5b_path}  (output folder: {week_tag})")

    print("\nNext step: run 05_export.py")


if __name__ == "__main__":
    main()
